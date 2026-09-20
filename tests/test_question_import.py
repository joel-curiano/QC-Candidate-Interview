from io import BytesIO

import pytest
from openpyxl import load_workbook

from question_import import QuestionImportError, export_questions_bytes, parse_questions, template_bytes


def test_template_is_a_parseable_blank_workbook():
    workbook = load_workbook(BytesIO(template_bytes()))
    sheet = workbook['Questions']
    sheet.append(['Welding QC', 'mcq', 'Choose one', 'A\nB', 'A', ''])
    sheet.append(['Welding QC', 'essay', 'Explain the process', '', '', 'Award for evidence.'])
    sheet.append(['Welding QC', 'oral_practical', 'Demonstrate the inspection process', '', '', 'Award for safe execution.'])
    sheet.append(['Welding QC', 'oral_practical', 'Explain the inspection decision', '', '', 'Award for technical accuracy.'])
    sheet.append(['Welding QC', 'oral_practical', 'Perform the inspection steps', '', '', 'Award for correct execution.'])
    output = BytesIO()
    workbook.save(output)
    results = parse_questions(output.getvalue())
    assert [result['success'] for result in results] == [True] * 5
    assert [result['question']['kind'] for result in results] == ['mcq', 'essay', 'oral_practical', 'oral_practical', 'oral_practical']
    assert results[0]['question']['options'] == ['A', 'B']


def test_template_options_keep_semicolons_inside_newline_delimited_options():
    workbook = load_workbook(BytesIO(template_bytes()))
    workbook['Questions'].append([
        'Welding QC', 'mcq', 'Choose one', 'Option A; with detail\nOption B', 'Option A; with detail', '',
    ])
    output = BytesIO()
    workbook.save(output)

    result = parse_questions(output.getvalue())[0]
    assert result['success'] is True
    assert result['question']['options'] == ['Option A; with detail', 'Option B']


def test_import_rejects_invalid_rows():
    workbook = load_workbook(BytesIO(template_bytes()))
    workbook['Questions'].append(['Welding QC', 'mcq', 'Choose one', 'A\nA', 'A', ''])
    output = BytesIO()
    workbook.save(output)
    results = parse_questions(output.getvalue())
    assert len(results) == 1
    assert results[0]['success'] is False
    assert 'unique options' in results[0]['error']

def test_existing_questions_export_is_reimportable():
    questions = [
        {
            'discipline': 'Welding QC',
            'q_type': 'mcq',
            'question_text': 'Choose one',
            'options': '["A", "B"]',
            'correct_answer': 'A',
            'rubric': '',
        },
        {
            'discipline': 'Civil QC',
            'q_type': 'essay',
            'question_text': 'Explain the inspection',
            'options': None,
            'correct_answer': None,
            'rubric': 'Award for evidence.',
        },
    ]

    exported = export_questions_bytes(questions)
    parsed = parse_questions(exported)

    assert [result['success'] for result in parsed] == [True, True]
    assert parsed[0]['question']['options'] == ['A', 'B']
    assert parsed[1]['question']['rubric'] == 'Award for evidence.'
