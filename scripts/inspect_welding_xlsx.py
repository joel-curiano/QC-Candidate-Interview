import openpyxl

wb = openpyxl.load_workbook('qc-question-template-Welding.xlsx')
print("Sheetnames:", wb.sheetnames)
s = wb.active
print(f"Rows: {s.max_row}, Cols: {s.max_column}")
for r in range(1, min(10, s.max_row + 1)):
    print([s.cell(r, c).value for c in range(1, s.max_column + 1)])
