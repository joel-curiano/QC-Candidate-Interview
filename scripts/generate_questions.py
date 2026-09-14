#!/usr/bin/env python3
"""
ARAMCO-Aligned QC Question Generator
Clears all existing question rows from qc-question-template.xlsx and repopulates
with 100 rigorous questions per discipline (40 MCQ + 20 Essay + 20 Oral + 20 Practical).
Total: 1,100 questions across 11 disciplines.

Run from project root:  python scripts/generate_questions.py
"""

import os
import sys

# Make sure imports resolve from project root
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from copy import copy

TEMPLATE_PATH = os.path.join(ROOT, "qc-question-template.xlsx")

# ── Import each discipline's question list ──────────────────────────────────
from scripts.questions.welding_qc      import QUESTIONS as WLD_Q
from scripts.questions.piping_qc       import QUESTIONS as PIP_Q
from scripts.questions.ndt_qc          import QUESTIONS as NDT_Q
from scripts.questions.mechanical_qc   import QUESTIONS as MEC_Q
from scripts.questions.electrical_qc   import QUESTIONS as ELC_Q
from scripts.questions.instrumentation_qc import QUESTIONS as INS_Q
from scripts.questions.e_and_i_qc      import QUESTIONS as EAI_Q
from scripts.questions.coating_qc      import QUESTIONS as COA_Q
from scripts.questions.civil_qc        import QUESTIONS as CIV_Q
from scripts.questions.cp_qc           import QUESTIONS as CAT_Q
from scripts.questions.comm_qc         import QUESTIONS as COM_Q

ALL_QUESTIONS = (
    WLD_Q + PIP_Q + NDT_Q + MEC_Q + ELC_Q +
    INS_Q + EAI_Q + COA_Q + CIV_Q + CAT_Q + COM_Q
)


def validate_questions(questions):
    """Basic structural validation before writing."""
    errors = []
    for i, q in enumerate(questions, 1):
        if len(q) != 7:
            errors.append(f"Row {i}: expected 7 fields, got {len(q)}")
            continue
        disc, kind, text, opts, correct, rubric, pts = q
        if kind not in ("mcq", "essay", "oral", "practical"):
            errors.append(f"Row {i}: invalid kind '{kind}'")
        if kind == "mcq":
            option_list = [o.strip() for o in opts.replace(";", "\n").splitlines() if o.strip()]
            if len(option_list) < 4:
                errors.append(f"Row {i} MCQ: fewer than 4 options")
            
            found = False
            for opt in option_list:
                if correct.strip().lower() in opt.lower() or opt.lower() in correct.strip().lower():
                    found = True
                    break

            if not found:
                errors.append(f"Row {i} MCQ: correct answer '{correct}' not found in options {option_list}")
            
            if pts != 1:
                errors.append(f"Row {i} MCQ: max_points must be 1, got {pts}")
        else:
            if not rubric:
                errors.append(f"Row {i} {kind}: rubric is empty")
    return errors


def clear_and_write(template_path, questions):
    wb = load_workbook(template_path)
    ws = wb["Questions"]

    # Delete all data rows (keep row 1 = headers)
    if ws.max_row and ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)

    # Write new questions
    for q in questions:
        disc, kind, text, opts, correct, rubric, pts = q
        ws.append([disc, kind, text, opts, correct, rubric, pts])

    wb.save(template_path)
    return ws.max_row - 1  # number of data rows written


def print_summary(questions):
    from collections import Counter
    by_disc = Counter()
    by_kind = Counter()
    for q in questions:
        by_disc[q[0]] += 1
        by_kind[(q[0], q[1])] += 1

    print("\n=== Question Generation Summary ===")
    for disc in sorted(set(q[0] for q in questions)):
        mcq  = by_kind[(disc, "mcq")]
        ess  = by_kind[(disc, "essay")]
        oral = by_kind[(disc, "oral")]
        prac = by_kind[(disc, "practical")]
        total = mcq + ess + oral + prac
        status = "✓" if (mcq == 40 and ess == 20 and oral == 20 and prac == 20) else "✗"
        print(f"  {status} {disc:35s} MCQ:{mcq:3d} Essay:{ess:3d} Oral:{oral:3d} Practical:{prac:3d} | Total:{total:4d}")

    print(f"\n  Grand Total: {len(questions)} questions across {len(set(q[0] for q in questions))} disciplines")


if __name__ == "__main__":
    print(f"Loading template: {TEMPLATE_PATH}")

    # Validate
    errors = validate_questions(ALL_QUESTIONS)
    if errors:
        print(f"\n[VALIDATION ERRORS — fix before writing]\n")
        for e in errors:
            print(f"  • {e}")
        sys.exit(1)
    print(f"Validation passed for {len(ALL_QUESTIONS)} questions.")

    # Write
    written = clear_and_write(TEMPLATE_PATH, ALL_QUESTIONS)
    print(f"Written {written} rows to {TEMPLATE_PATH}")

    print_summary(ALL_QUESTIONS)
    print("\nDone. Template is ready for import.\n")
