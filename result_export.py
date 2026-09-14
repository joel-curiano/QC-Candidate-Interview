"""Excel export for the assessment results log."""
from io import BytesIO
import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
import database as db

ILLEGAL_EXCEL_CHARS = re.compile(r'[\x00-\x08\x0B\x0C\x0E-\x1F]')


RESULT_HEADERS = [
    'Reference', 'Candidate', 'Job Title', 'Employee No', 'Discipline', 'Project Location', 'Project Assignment', 'Exam Date',
    'Status', 'Multiple Choice Grade', 'Essay Grade', 'Oral Grade', 'Practical Grade',
    'Reviewer Comments', 'Graded (UTC)', 'Overall Result',
]


def _safe_cell(value):
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.replace(tzinfo=None)
    if isinstance(value, str):
        value = ILLEGAL_EXCEL_CHARS.sub('', value)
        if value.lstrip().startswith(('=', '+', '-', '@')):
            return "'" + value
    return '' if value is None else value


def _row_values(row):
    return [
        row.get('id'), row.get('candidate_name'), row.get('designation', ''), row.get('employee_no', ''),
        row.get('discipline', ''), row.get('project_location', ''), row.get('project_assignment', row.get('project_location', '')),
        row.get('exam_date', ''), row.get('status', ''),
        row.get('multiple_choice_grade', db.category_result(row, 'mcq')), row.get('essay_grade', db.category_result(row, 'essay')),
        row.get('oral_grade', db.category_result(row, 'oral')), row.get('practical_grade', db.category_result(row, 'practical')),
        row.get('reviewer_comments', ''), row.get('graded_at', ''),
        row.get('result', db.result(row)),
    ]


def excel_bytes(rows):
    """Return a formatted Excel workbook containing the results log."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'CTA Record Log'
    sheet.append(RESULT_HEADERS)
    for row in rows:
        sheet.append([_safe_cell(value) for value in _row_values(row)])

    header_fill = PatternFill('solid', fgColor='1F4E78')
    for cell in sheet[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=cell.column in (2, 3, 5, 8, 9, 21))

    widths = [12, 24, 30, 18, 20, 18, 18, 18, 24, 20, 16, 22, 18, 20, 16, 16, 16, 16, 18, 18, 42, 22]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = sheet.dimensions
    sheet.row_dimensions[1].height = 30

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()
